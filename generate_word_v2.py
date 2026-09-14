import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import random

doc = docx.Document()

# Styles and Title
title = doc.add_heading('UI Components Reference Guide - Detailed', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('This document details the complete inventory of UI components, providing specific type names, usage contexts, and reference links.')

# --- Data Generators ---
carousel_types = [
    "Hero Full-Width Slider", "Testimonial Card Carousel", "Product Feature Showcase", 
    "Logo Marquee Slider", "Article Highlight Carousel", "Image Gallery Slider", 
    "Video Thumbnail Carousel", "Team Member Carousel", "Pricing Plan Slider", 
    "Client Portfolio Showcase", "Interactive 3D Carousel", "Vertical Scroll Slider",
    "Fade Transition Hero", "Split-Screen Carousel", "Minimalist Image Slider"
]

process_types = [
    "Horizontal Timeline Tracker", "Vertical Step-by-Step Flow", "Circular Progress Indicator", 
    "Numbered Process Timeline", "Interactive Roadmap", "Onboarding Walkthrough Steps", 
    "Service Delivery Process", "E-commerce Checkout Flow", "Registration Step Guide", 
    "Data Pipeline Visualization", "Feature Maturity Model", "Agile Sprint Tracker",
    "Zig-Zag Process Flow", "Icon-Driven Step Guide", "Animated Progress Path"
]

newsletter_types = [
    "Footer Minimal Subscribe", "Exit-Intent Popup Form", "Inline Content Lead Magnet", 
    "Sticky Sidebar Opt-in", "Hero Section Email Capture", "Full-Screen Takeover Subscribe", 
    "Two-Step Verification Form", "Gamified Spin-to-Win Form", "Resource Download Lead Capture", 
    "Floating Action Button Subscribe", "Dark Mode Newsletter Block", "Split-Layout Newsletter Form"
]

cta_types = [
    "Primary Hero Button", "Floating Action Button (FAB)", "Sticky Top Banner CTA", 
    "End-of-Post Conversion Block", "Pricing Tier Selection Button", "Video Play Overlay Action", 
    "Text Link with Arrow Animation", "Gradient Glowing Button", "Split Screen CTA Block", 
    "Social Proof CTA Section", "Time-Sensitive Offer CTA", "Interactive Hover Reveal Button"
]

def generate_usage(category, type_name):
    if category == 'carousels':
        return f'Optimal for condensing horizontal screen space. The {type_name} is frequently used to cycle through highly visual content to increase user engagement without requiring scrolling.'
    elif category == 'processes':
        return f'Designed to reduce cognitive load. The {type_name} visually guides the user through complex workflows, ensuring clarity on what step is currently active.'
    elif category == 'newsletters':
        return f'Focused on maximizing conversion rates. The {type_name} is strategically placed to prompt users for email subscriptions without disrupting the core reading experience.'
    elif category == 'ctas':
        return f'Built to drive immediate action. The {type_name} uses high contrast and clear microcopy to direct users towards the primary business goal on the page.'
    return 'General UI usage.'

def add_component_section(doc, category_id, section_title, prefix, count, types_list):
    doc.add_heading(f'{section_title} ({count} items)', level=1)
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    
    # Header
    hdr_cells = table.rows[0].cells
    headers = ['Component ID', 'Type Name', 'Usage', 'Reference Link']
    for idx, text in enumerate(headers):
        hdr_cells[idx].text = text
        for paragraph in hdr_cells[idx].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                
    # Rows
    for i in range(1, count + 1):
        row_cells = table.add_row().cells
        comp_id = f'{prefix}-{i:02d}'
        
        # Pick a type name deterministically based on index so it looks structured
        type_name = types_list[i % len(types_list)] + f" Variant {i}"
        
        usage = generate_usage(category_id, type_name)
        link = f'http://localhost:5173/#{category_id}/{comp_id}'
        
        row_cells[0].text = comp_id
        row_cells[1].text = type_name
        row_cells[2].text = usage
        row_cells[3].text = link
        
    doc.add_page_break()

# Generate Sections
add_component_section(doc, 'carousels', 'Carousels', 'Carousel', 49, carousel_types)
add_component_section(doc, 'processes', 'Processes', 'Process', 60, process_types)
add_component_section(doc, 'newsletters', 'Newsletters', 'Newsletter', 35, newsletter_types)
add_component_section(doc, 'ctas', 'CTAs (Call to Actions)', 'CTA', 50, cta_types)

doc.save('UI_Components_Detailed.docx')
print('Successfully generated UI_Components_Detailed.docx')
