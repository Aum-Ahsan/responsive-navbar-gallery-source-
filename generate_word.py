import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Title
title = doc.add_heading('UI Components Reference Guide', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph(
    'This document provides a comprehensive list of the UI components available in the project, '
    'including their names and typical usages.'
)

def add_component_section(doc, section_title, prefix, count, description_template):
    # Section Heading
    doc.add_heading(f'{section_title} ({count} items)', level=1)
    
    # Create Table
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    # Set column widths (approximate, Word handles it dynamically but this helps)
    table.autofit = False
    table.columns[0].width = Inches(0.8)
    table.columns[1].width = Inches(2.0)
    table.columns[2].width = Inches(3.5)
    
    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'S.No'
    hdr_cells[1].text = 'Component Name'
    hdr_cells[2].text = 'Usage / Description'
    
    # Bold headers
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                
    # Add Rows
    for i in range(1, count + 1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(i)
        
        # Name
        name = f'{prefix}-{i:02d}'
        row_cells[1].text = name
        
        # Usage
        # Make usage descriptions slightly varied or just use a good professional template
        usage = description_template.format(index=i)
        row_cells[2].text = usage
        
    doc.add_page_break()

# Add Carousels
add_component_section(
    doc,
    'Carousels',
    'Carousel',
    49,
    'A rotating image or content gallery. Carousel variant {index} is ideal for showcasing multiple items (like testimonials, products, or feature highlights) within a limited horizontal space, keeping users engaged without overwhelming the page layout.'
)

# Add Processes
add_component_section(
    doc,
    'Processes',
    'Process',
    60,
    'A step-by-step workflow or timeline component. Process variant {index} is designed to visually guide the user through a sequence of steps, stages, or instructions, making complex workflows easy to understand at a glance.'
)

# Add Newsletters
add_component_section(
    doc,
    'Newsletters',
    'Newsletter',
    35,
    'An email capture or subscription form section. Newsletter variant {index} offers a distinct layout to attract user attention and encourage sign-ups, helping to build a mailing list effectively with optimized form placement.'
)

# Add CTAs
add_component_section(
    doc,
    'CTAs (Call to Actions)',
    'CTA',
    50,
    'A prominent section designed to drive user action. CTA variant {index} leverages contrast, spacing, and compelling typography to guide users toward key conversion points like "Sign Up", "Buy Now", or "Learn More".'
)

# Save Document
doc.save('UI_Components_List.docx')
print('Successfully generated UI_Components_List.docx')
